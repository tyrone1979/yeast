#!/usr/bin/env python3
"""Convert industrial yeast fermentation Excel workbook into publication-ready CSVs."""

from __future__ import annotations

import csv
import json
import math
import statistics as stats
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "data" / "data.xlsx"
OUT = ROOT / "dataset"
ASSESS = "评价标准"


def round_or_none(v, nd=6):
    if v is None:
        return ""
    if isinstance(v, float):
        if math.isnan(v):
            return ""
        return round(v, nd)
    return v


def to_float(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def describe(rows, key, batch_ids=None):
    vals = []
    for r in rows:
        if batch_ids is not None and r["batch_id"] not in batch_ids:
            continue
        v = to_float(r.get(key))
        if v is not None:
            vals.append(v)
    if not vals:
        return None
    return {
        "n": len(vals),
        "mean": round(sum(vals) / len(vals), 4),
        "std": round(stats.pstdev(vals), 4),
        "min": round(min(vals), 4),
        "max": round(max(vals), 4),
    }


def write_csv(path: Path, fields, rows):
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "figures").mkdir(parents=True, exist_ok=True)

    wb = load_workbook(SRC, data_only=True)

    meta_fields = [
        "batch_id",
        "lot_no",
        "strain_id",
        "recipe",
        "seed_lot_no",
        "seed_volume_m3",
        "seed_y30_kg",
        "yield_y30_kg",
        "n_records",
        "time_start_h",
        "time_end_h",
        "data_completeness",
        "has_process_variables",
        "notes",
    ]
    ts_fields = [
        "batch_id",
        "lot_no",
        "time_h",
        "airflow_m3_h",
        "volume_m3",
        "ph",
        "alcohol_vv_pct",
        "cell_concentration_gpl",
        "biomass_y30_kg",
        "growth_modulus",
        "sugar_feed_rate_kg_rs_h",
        "total_fermentable_sugar_kg",
        "remark",
        "delta_alcohol",
        "delta_growth_modulus",
        "auto_growth_modulus",
    ]

    meta_rows = []
    ts_rows = []

    for name in wb.sheetnames:
        if name in ("DIFF", ASSESS):
            continue
        ws = wb[name]
        lot = ws.cell(2, 2).value
        seed_no = ws.cell(2, 5).value
        yield_y30 = ws.cell(2, 8).value
        strain = ws.cell(3, 2).value
        seed_m3 = ws.cell(3, 5).value
        recipe = ws.cell(4, 2).value
        seed_y30 = ws.cell(4, 5).value

        records = []
        for r in range(6, ws.max_row + 1):
            t = ws.cell(r, 1).value
            if t is None:
                continue
            vals = [ws.cell(r, c).value for c in range(1, 15)]
            if all(v is None for v in vals[1:]):
                continue
            records.append(vals)

        air_n = sum(1 for v in records if v[1] is not None)
        complete = "full_process" if air_n >= 10 else "alcohol_sugar_biomass"
        notes = []
        if name == "05":
            notes.append(
                "volume_m3 values appear ~10x larger than peer batches; retained as recorded"
            )
        if lot == 2616:
            notes.append("lot_no 2616 appears in more than one sheet")
        if yield_y30 is None:
            notes.append("yield_y30_kg missing in source sheet header")

        times = [v[0] for v in records if isinstance(v[0], (int, float))]
        batch_id = f"B{name}"
        meta_rows.append(
            {
                "batch_id": batch_id,
                "lot_no": lot if lot is not None else "",
                "strain_id": strain if strain is not None else "",
                "recipe": recipe if recipe is not None else "",
                "seed_lot_no": seed_no if seed_no is not None else "",
                "seed_volume_m3": round_or_none(seed_m3, 4),
                "seed_y30_kg": round_or_none(seed_y30, 4),
                "yield_y30_kg": round_or_none(yield_y30, 4),
                "n_records": len(records),
                "time_start_h": min(times) if times else "",
                "time_end_h": max(times) if times else "",
                "data_completeness": complete,
                "has_process_variables": "yes" if complete == "full_process" else "no",
                "notes": "; ".join(notes),
            }
        )

        for vals in records:
            ts_rows.append(
                {
                    "batch_id": batch_id,
                    "lot_no": lot if lot is not None else "",
                    "time_h": vals[0],
                    "airflow_m3_h": round_or_none(vals[1], 3),
                    "volume_m3": round_or_none(vals[2], 4),
                    "ph": round_or_none(vals[3], 4),
                    "alcohol_vv_pct": round_or_none(vals[4], 6),
                    "cell_concentration_gpl": round_or_none(vals[5], 4),
                    "biomass_y30_kg": round_or_none(vals[6], 4),
                    "growth_modulus": round_or_none(vals[7], 6),
                    "sugar_feed_rate_kg_rs_h": round_or_none(vals[8], 4),
                    "total_fermentable_sugar_kg": round_or_none(vals[9], 4),
                    "remark": vals[10] if vals[10] is not None else "",
                    "delta_alcohol": round_or_none(vals[11], 6),
                    "delta_growth_modulus": round_or_none(vals[12], 6),
                    "auto_growth_modulus": round_or_none(vals[13], 6),
                }
            )

    write_csv(OUT / "batch_metadata.csv", meta_fields, meta_rows)
    write_csv(OUT / "fermentation_timeseries.csv", ts_fields, ts_rows)

    # Evaluation / quality profiles
    ws = wb[ASSESS]
    prof = []
    for r in range(6, ws.max_row + 1):
        t = ws.cell(r, 12).value
        if t is None:
            continue
        prof.append(
            {
                "time_h": t,
                "alcohol_profile_vv_pct": round_or_none(ws.cell(r, 13).value, 6),
                "growth_modulus_profile": round_or_none(ws.cell(r, 14).value, 6),
            }
        )
    write_csv(
        OUT / "evaluation_profiles.csv",
        ["time_h", "alcohol_profile_vv_pct", "growth_modulus_profile"],
        prof,
    )

    # Protocol hourly setpoints
    ws = wb["DIFF"]
    sp_fields = [
        "brew_hour",
        "hour_interval",
        "temp_c",
        "airflow_m3_h",
        "ph_setpoint",
        "volume_m3",
        "fill_pct",
        "alcohol_target_vv_pct",
        "cell_apparent_gpl",
        "cell_actual_y30_gpl",
        "biomass_y30_kg",
        "growth_modulus",
        "sugar_feed_rate_kg_rs_h",
        "total_rs_kg",
        "wort_kg",
        "ammonia_rate_kg_h",
        "total_ammonia_kg",
        "phosphoric_rate_kg_h",
        "total_h3po4_kg",
        "note",
    ]
    setpoints = []
    for r in range(33, 50):
        hour_label = ws.cell(r, 2).value
        if hour_label is None:
            continue
        if isinstance(hour_label, (int, float)):
            brew_hour = int(hour_label)
        else:
            s = str(hour_label).strip().lower()
            if "cool" in s:
                break
            if "start" in s:
                brew_hour = 0
            else:
                try:
                    brew_hour = int(float(s))
                except ValueError:
                    continue
        setpoints.append(
            {
                "brew_hour": brew_hour,
                "hour_interval": ws.cell(r, 3).value
                if ws.cell(r, 3).value is not None
                else "",
                "temp_c": round_or_none(ws.cell(r, 4).value, 3),
                "airflow_m3_h": round_or_none(ws.cell(r, 5).value, 3),
                "ph_setpoint": round_or_none(ws.cell(r, 6).value, 3),
                "volume_m3": round_or_none(ws.cell(r, 8).value, 4),
                "fill_pct": round_or_none(ws.cell(r, 9).value, 4),
                "alcohol_target_vv_pct": round_or_none(ws.cell(r, 10).value, 6),
                "cell_apparent_gpl": round_or_none(ws.cell(r, 11).value, 4),
                "cell_actual_y30_gpl": round_or_none(ws.cell(r, 12).value, 4),
                "biomass_y30_kg": round_or_none(ws.cell(r, 13).value, 4),
                "growth_modulus": round_or_none(ws.cell(r, 14).value, 6),
                "sugar_feed_rate_kg_rs_h": round_or_none(ws.cell(r, 15).value, 4),
                "total_rs_kg": round_or_none(ws.cell(r, 16).value, 4),
                "wort_kg": round_or_none(ws.cell(r, 17).value, 4),
                "ammonia_rate_kg_h": round_or_none(ws.cell(r, 18).value, 4),
                "total_ammonia_kg": round_or_none(ws.cell(r, 19).value, 4),
                "phosphoric_rate_kg_h": round_or_none(ws.cell(r, 20).value, 4),
                "total_h3po4_kg": round_or_none(ws.cell(r, 21).value, 4),
                "note": ws.cell(r, 22).value if ws.cell(r, 22).value is not None else "",
            }
        )
    write_csv(OUT / "protocol_hourly_setpoints.csv", sp_fields, setpoints)

    materials = [
        {
            "item": "protocol_name",
            "value": "ZHENAO TRIAL DIFFERENTIAL BREW PROTOCOL-NOV.2014",
            "unit": "",
            "category": "protocol",
        },
        {
            "item": "protocol_date",
            "value": "2014-11-18",
            "unit": "",
            "category": "protocol",
        },
        {
            "item": "target_protein",
            "value": "57-58",
            "unit": "%",
            "category": "product_target",
        },
        {
            "item": "target_p2o5",
            "value": "2.8-3.0",
            "unit": "%",
            "category": "product_target",
        },
        {
            "item": "assumed_yield_y30_per_rs",
            "value": "1.55",
            "unit": "kg Y30 / kg RS",
            "category": "assumption",
        },
        {
            "item": "assumed_naf",
            "value": "1.0-1.05",
            "unit": "",
            "category": "assumption",
        },
        {
            "item": "assumed_paf",
            "value": "0.85-0.95",
            "unit": "",
            "category": "assumption",
        },
        {
            "item": "wort_source",
            "value": "Beet 100% / Cane 0%",
            "unit": "",
            "category": "recipe",
        },
        {
            "item": "average_wort_concentration",
            "value": "0.35",
            "unit": "kg RS / kg wort",
            "category": "recipe",
        },
        {
            "item": "wort_density",
            "value": "1.24",
            "unit": "kg/L",
            "category": "recipe",
        },
        {"item": "total_rs", "value": "15000", "unit": "kg", "category": "materials"},
        {
            "item": "total_wort",
            "value": "42857.14",
            "unit": "kg",
            "category": "materials",
        },
        {
            "item": "ammonia_100pct",
            "value": "800",
            "unit": "kg",
            "category": "materials",
        },
        {
            "item": "h3po4_100pct",
            "value": "420",
            "unit": "kg",
            "category": "materials",
        },
        {"item": "mgso4", "value": "100", "unit": "kg", "category": "materials"},
        {"item": "znso4", "value": "5", "unit": "kg", "category": "materials"},
        {"item": "cuso4", "value": "180", "unit": "g", "category": "materials"},
        {"item": "thiamine_b1", "value": "1500", "unit": "g", "category": "materials"},
        {
            "item": "ca_pantothenate_b5",
            "value": "2500",
            "unit": "g",
            "category": "materials",
        },
        {"item": "pyridoxine_b6", "value": "800", "unit": "g", "category": "materials"},
        {"item": "biotin", "value": "25", "unit": "g", "category": "materials"},
        {
            "item": "fermenter_diameter",
            "value": "4",
            "unit": "m",
            "category": "equipment",
        },
        {
            "item": "fermenter_height",
            "value": "12",
            "unit": "m",
            "category": "equipment",
        },
        {
            "item": "nominal_fermenter_volume",
            "value": "151.30",
            "unit": "m3",
            "category": "equipment",
        },
        {
            "item": "fermentation_time",
            "value": "16",
            "unit": "h",
            "category": "process",
        },
        {
            "item": "maturation_time",
            "value": "40",
            "unit": "min",
            "category": "process",
        },
        {
            "item": "seed_transfer_time",
            "value": "20",
            "unit": "min",
            "category": "process",
        },
        {
            "item": "projected_cream_output",
            "value": "27250",
            "unit": "kg Y30",
            "category": "process",
        },
    ]
    write_csv(
        OUT / "recipe_parameters.csv",
        ["item", "value", "unit", "category"],
        materials,
    )

    full_ids = {m["batch_id"] for m in meta_rows if m["data_completeness"] == "full_process"}
    summary = {
        "n_batches": len(meta_rows),
        "n_full_process_batches": len(full_ids),
        "n_sparse_batches": len(meta_rows) - len(full_ids),
        "n_timeseries_rows": len(ts_rows),
        "strain_ids": sorted({str(m["strain_id"]) for m in meta_rows if m["strain_id"] != ""}),
        "recipes": sorted({str(m["recipe"]) for m in meta_rows if m["recipe"] != ""}),
        "lot_no_min": min(int(m["lot_no"]) for m in meta_rows if m["lot_no"] != ""),
        "lot_no_max": max(int(m["lot_no"]) for m in meta_rows if m["lot_no"] != ""),
        "records_per_batch": sorted({m["n_records"] for m in meta_rows}),
        "descriptive_all": {
            k: describe(ts_rows, k)
            for k in [
                "airflow_m3_h",
                "volume_m3",
                "ph",
                "alcohol_vv_pct",
                "cell_concentration_gpl",
                "biomass_y30_kg",
                "growth_modulus",
                "sugar_feed_rate_kg_rs_h",
                "total_fermentable_sugar_kg",
            ]
        },
        "descriptive_full_process": {
            k: describe(ts_rows, k, full_ids)
            for k in [
                "airflow_m3_h",
                "volume_m3",
                "ph",
                "alcohol_vv_pct",
                "cell_concentration_gpl",
                "biomass_y30_kg",
                "growth_modulus",
                "sugar_feed_rate_kg_rs_h",
                "total_fermentable_sugar_kg",
            ]
        },
        "yield_stats": describe(
            [
                {
                    "batch_id": m["batch_id"],
                    "yield_y30_kg": m["yield_y30_kg"],
                }
                for m in meta_rows
            ],
            "yield_y30_kg",
        ),
    }
    (OUT / "summary_stats.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8"
    )

    print(f"Wrote {len(meta_rows)} batches, {len(ts_rows)} timeseries rows to {OUT}")
    print(f"Full-process batches: {sorted(full_ids)}")


if __name__ == "__main__":
    main()
