#!/usr/bin/env python3
"""Convert the Data in Brief markdown draft to a clean DOCX with embedded figures."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "paper" / "yeast_fermentation_dataset_DataInBrief.md"
OUT = ROOT / "paper" / "yeast_fermentation_dataset_DataInBrief.docx"
FIG = ROOT / "paper" / "figures"

FIGURE_FILES = {
    1: FIG / "fig1_sugar_feed_trajectories.png",
    2: FIG / "fig2_alcohol_biomass_b01.png",
    3: FIG / "fig3_mean_ph_airflow.png",
    4: FIG / "fig4_sugar_feed_hist.png",
}


def set_run_font(run, size=11, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_hyperlink(paragraph, text: str, url: str, size: int = 11):
    """Add a clickable hyperlink run to a paragraph."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rPr.append(sz)
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "Times New Roman")
    rPr.append(rFonts)
    new_run.append(rPr)
    text_elem = OxmlElement("w:t")
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def clean_md_text(text: str) -> str:
    """Remove markdown/latex leftovers for plain Word text (no link expansion)."""
    text = text.replace("$^{1,*}$", "¹,*")
    text = text.replace("$^{1}$", "¹")
    text = text.replace("$^{1,}$", "¹")
    text = re.sub(r"\$\^\{([^}]+)\}\$", r"\1", text)
    text = text.replace("$", "")
    # strip inline code backticks but keep content
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = text.replace("****", "")
    return text


def _is_url(s: str) -> bool:
    return s.startswith("http://") or s.startswith("https://")


def add_runs_with_markup(paragraph, text: str, size=11):
    text = clean_md_text(text)
    # Split out markdown links and bare URLs so DOCX gets proper hyperlinks.
    link_re = re.compile(
        r"\[([^\]]+)\]\(([^)]+)\)|(https?://[^\s\]\)]+)"
    )
    pos = 0
    for m in link_re.finditer(text):
        if m.start() > pos:
            _add_markup_fragment(paragraph, text[pos : m.start()], size)
        if m.group(1) is not None:
            label, url = m.group(1), m.group(2)
            # Prefer canonical DOI URL as visible text when label is a bare DOI.
            if label.startswith("10.") and "doi.org/" in url:
                display = url
            elif _is_url(label):
                display = label
            else:
                display = label
            add_hyperlink(paragraph, display, url, size=size)
        else:
            url = m.group(3).rstrip(".,;")
            add_hyperlink(paragraph, url, url, size=size)
            # re-append trailing punctuation stripped from URL if present
            trail = m.group(3)[len(url) :]
            if trail:
                run = paragraph.add_run(trail)
                set_run_font(run, size=size)
        pos = m.end()
    if pos < len(text):
        _add_markup_fragment(paragraph, text[pos:], size)


def _add_markup_fragment(paragraph, text: str, size: int = 11):
    # bold segments **...**
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True)
        else:
            subparts = re.split(r"(\*[^*]+\*)", part)
            for sp in subparts:
                if not sp:
                    continue
                if sp.startswith("*") and sp.endswith("*") and not sp.startswith("**"):
                    run = paragraph.add_run(sp[1:-1])
                    set_run_font(run, size=size, italic=True)
                else:
                    run = paragraph.add_run(sp)
                    set_run_font(run, size=size)


def add_table_from_lines(doc, lines):
    rows = []
    for line in lines:
        if not line.strip().startswith("|"):
            continue
        cells = [clean_md_text(c.strip()) for c in line.strip().strip("|").split("|")]
        if all(set(c) <= set("-: ") for c in cells):
            continue
        rows.append(cells)
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            add_runs_with_markup(p, val, size=10)
            for run in p.runs:
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


def add_figure(doc, fig_no: int, caption: str):
    path = FIGURE_FILES.get(fig_no)
    if path and path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(path), width=Inches(5.9))
    else:
        p = doc.add_paragraph()
        run = p.add_run(f"[Missing figure file for Fig. {fig_no}]")
        set_run_font(run, italic=True)
        run.font.color.rgb = RGBColor(180, 0, 0)

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs_with_markup(cap, caption, size=10)
    for run in cap.runs:
        run.italic = True
    doc.add_paragraph()


def main():
    text = MD.read_text(encoding="utf-8")
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    pf = style.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE

    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue

        if line.startswith("# "):
            p = doc.add_heading(clean_md_text(line[2:].strip()), level=1)
            for run in p.runs:
                set_run_font(run, size=14, bold=True)
            i += 1
            continue

        if line.startswith("## "):
            p = doc.add_heading(clean_md_text(line[3:].strip()), level=2)
            for run in p.runs:
                set_run_font(run, size=12, bold=True)
            i += 1
            continue

        # markdown image
        m_img = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line.strip())
        if m_img:
            # caption may follow; image path relative
            alt = m_img.group(1)
            rel = m_img.group(2)
            img_path = (MD.parent / rel).resolve()
            if not img_path.exists():
                img_path = (ROOT / rel.lstrip("./")).resolve()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if img_path.exists():
                p.add_run().add_picture(str(img_path), width=Inches(5.9))
            else:
                run = p.add_run(f"[Missing image: {rel}]")
                set_run_font(run, italic=True)
            if alt:
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_runs_with_markup(cap, alt, size=10)
                for run in cap.runs:
                    run.italic = True
            i += 1
            continue

        if line.strip().startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table_from_lines(doc, block)
            continue

        if line.strip().startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                i += 1
            i += 1
            continue

        # Figure caption lines: embed PNG then caption
        m_fig = re.match(r"\*\*Fig\.\s*(\d+)\.\*\*\s*(.*)", line.strip())
        if m_fig:
            fig_no = int(m_fig.group(1))
            caption_body = clean_md_text(m_fig.group(2))
            # drop parenthetical file paths from caption
            caption_body = re.sub(r"\s*\([^)]*figures/[^)]*\)\s*", "", caption_body).strip()
            caption = f"Fig. {fig_no}. {caption_body}"
            add_figure(doc, fig_no, caption)
            i += 1
            continue

        if line.lstrip().startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs_with_markup(p, line.lstrip()[2:])
            i += 1
            continue

        if re.match(r"^\d+[\.)]\s+", line):
            p = doc.add_paragraph(style="List Number")
            add_runs_with_markup(p, re.sub(r"^\d+[\.)]\s+", "", line))
            i += 1
            continue

        # label-only bold line like **Authors:**
        if re.fullmatch(r"\*\*[^*]+\*\*", line.strip()):
            p = doc.add_paragraph()
            add_runs_with_markup(p, line.strip())
            i += 1
            continue

        p = doc.add_paragraph()
        add_runs_with_markup(p, line)
        i += 1

    doc.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    main()
